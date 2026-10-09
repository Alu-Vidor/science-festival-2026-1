"""Build six editable light slides, PDF/PNG exports and a separate host script.

Run from any directory: python build_materials.py --guide-source <JSON>
Requires python-pptx, python-docx, reportlab, PyMuPDF and LibreOffice.
The generated illustration is inserted unchanged; slide text and diagrams are native.
"""
import argparse
import json
import math
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement as DocxElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGB
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parent
W, H = 13.333333, 7.5
FONT = 'Arial'
P = dict(bg='FCFBF8', ink='233D40', muted='657B7C', teal='307A71', mint='E4F1E8',
         lilac='EBEAF7', purple='7773A8', yellow='F8EED8', ochre='A48136',
         blue='E7EFF5', steel='60879C', line='D6E2DC', white='FFFFFF', gray='DFE5E2')


def rgb(key):
    return RGBColor.from_string(P.get(key, key))


def shape(slide, x, y, w, h, fill='white', border=None, kind=Shape.ROUNDED_RECTANGLE):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid(); item.fill.fore_color.rgb = rgb(fill)
    if border:
        item.line.color.rgb = rgb(border); item.line.width = Pt(1)
    else:
        item.line.fill.background()
    # Override the theme's default shadow for clean, light diagram cards.
    item._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in item._element.xpath('p:style/a:effectRef'):
        effect.set('idx', '0')
    if kind == Shape.ROUNDED_RECTANGLE:
        item.adjustments[0] = .12
    return item


def circle(slide, x, y, d, fill='teal', border=None):
    return shape(slide, x, y, d, d, fill, border, Shape.OVAL)


def line(slide, x1, y1, x2, y2, color='teal', width=2, arrow=False):
    from pptx.enum.shapes import MSO_CONNECTOR
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = rgb(color); item.line.width = Pt(width)
    if arrow:
        tip = shape(slide, x2 - .055, y2 - .055, .11, .11, color, kind=Shape.ISOSCELES_TRIANGLE)
        tip.rotation = math.degrees(math.atan2(y2 - y1, x2 - x1)) + 90
    return item


def wrapped(message, size, width, bold=False):
    lines = []
    for raw in message.split('\n'):
        words = raw.split(); current = ''
        for word in words:
            trial = (current + ' ' + word).strip()
            if current and pdfmetrics.stringWidth(trial, 'MetricBold' if bold else 'Metric', size) > width * 72 - 3:
                lines.append(current); current = word
            else:
                current = trial
        lines.append(current)
    return lines


def text(slide, message, x, y, w, h=None, size=22, color='ink', bold=False, align='left'):
    lines = wrapped(message, size, w, bold)
    minimum = (len(lines) * size * 1.22 + 4) / 72
    if h is None:
        h = minimum
    if h + .03 < minimum:
        raise ValueError(f'Text does not fit: {message!r}, needs {minimum:.2f}, has {h:.2f}')
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = item.text_frame; tf.clear(); tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, value in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = value; para.font.name = FONT; para.font.size = Pt(size)
        para.font.bold = bold; para.font.color.rgb = rgb(color)
        para.space_before = para.space_after = Pt(0); para.line_spacing = Pt(size * 1.22)
        para.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    return item


def chip(slide, label, x, y, w, fill='mint', color='teal'):
    shape(slide, x, y, w, .38, fill)
    text(slide, label, x + .12, y + .08, w - .24, size=12, color=color, bold=True)


def base(prs, n, content):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('bg')
    circle(slide, .62, .34, .13, 'teal')
    text(slide, 'ОСНОВЫ ИИ', .86, .3, 3, size=12, color='teal', bold=True)
    chip(slide, 'ПОЗНАЁМ ЧЕРЕЗ ОПЫТ', 9.85, .24, 2.84)
    text(slide, content['title'], .65, .89, 12.05, size=39, bold=True)
    text(slide, content['subtitle'], .68, 1.75, 11.9, size=23, color='muted')
    line(slide, .65, 6.96, 12.68, 6.96, 'line', .8)
    text(slide, 'поИИграем?  ·  Фестиваль науки', .68, 7.12, 8, size=12, color='muted')
    text(slide, f'{n:02d} / 06', 11.62, 7.09, 1.02, size=14, color='teal', bold=True, align='right')
    slide.notes_slide.notes_text_frame.text = content['notes']
    return slide


def banner(slide, message, fill='mint', color='teal', y=6.18, size=20):
    shape(slide, .65, y, 12.03, .58, fill)
    text(slide, message, .87, y + .12, 11.57, size=size, color=color, bold=True, align='center')


def wheel(slide, x, y, dirty=False, wet=False):
    for dx in [0, .48]:
        shape(slide, x + dx, y, .3, .49, 'teal')
        for dy in [.09, .21, .33]:
            line(slide, x + dx + .06, y + dy, x + dx + .24, y + dy, 'mint', 2)
        if dirty:
            circle(slide, x + dx + .04, y + .07, .08, 'ochre')
    if wet:
        circle(slide, x + .85, y + .27, .12, 'steel')


def model(slide, x, y, d=.64, color='teal'):
    shape(slide, x, y, d, d, color)
    for i in range(3):
        v = .13 + i * .16
        line(slide, x + v, y - .11, x + v, y + .05, color, 2)
        line(slide, x + v, y + d - .05, x + v, y + d + .11, color, 2)
        line(slide, x - .11, y + v, x + .05, y + v, color, 2)
        line(slide, x + d - .05, y + v, x + d + .11, y + v, color, 2)
    circle(slide, x + .2, y + .2, .22, 'mint')


def person(slide, x, y, color='teal', scale=1):
    circle(slide, x + .075 * scale, y, .15 * scale, color)
    shape(slide, x, y + .19 * scale, .3 * scale, .33 * scale, color)


def home(slide, x, y, color='teal', scale=1):
    shape(slide, x, y + .19 * scale, .66 * scale, .45 * scale, color)
    shape(slide, x - .05 * scale, y, .76 * scale, .36 * scale, color, kind=Shape.ISOSCELES_TRIANGLE)
    shape(slide, x + .27 * scale, y + .39 * scale, .14 * scale, .25 * scale, 'mint', kind=Shape.RECTANGLE)


def build_slides():
    source = json.loads((ROOT / 'slide-content.json').read_text(encoding='utf-8'))
    assert len(source) == 6
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    prs.core_properties.title = 'Искусственный интеллект: данные, обучение и проверка'
    prs.core_properties.subject = 'Вводная об ИИ с примерами робота и агентной симуляции города'
    prs.core_properties.author = 'поИИграем? · Фестиваль науки'
    prs.core_properties.keywords = 'ИИ, данные, обучение, робот, агенты, симуляция'

    s = base(prs, 1, source[0])
    s.shapes.add_picture(str(ROOT / 'assets' / 'robot-and-city.png'), Inches(6.66), Inches(2.3), width=Inches(6.0))
    text(s, 'Машинное обучение\nнаходит закономерности\nпо примерам.', .7, 2.65, 5.7, size=30, bold=True)
    for i, (label, fill, color) in enumerate([('Распознать', 'mint', 'teal'), ('Предсказать', 'lilac', 'purple'), ('Выбрать действие', 'yellow', 'ochre')]):
        shape(s, .7, 4.42 + i * .52, 4.35, .42, fill)
        text(s, label, .9, 4.48 + i * .52, 3.95, size=17, color=color, bold=True)
    text(s, 'Один из подходов к созданию ИИ.', .73, 6.12, 4.7, size=17, color='muted')

    s = base(prs, 2, source[1])
    cards = [('Что наблюдаем', 'Покрытие\nи колёса до шага', 'mint', 'teal'),
             ('Проводим опыт', 'Смотрим,\nчто произошло', 'yellow', 'ochre'),
             ('Измеряем', 'Расход · проезд\nКолёса после шага', 'blue', 'steel'),
             ('Передаём данные', 'Модель ищет\nзакономерности', 'lilac', 'purple')]
    for i, (heading, body, fill, color) in enumerate(cards):
        x = .65 + i * 3.08
        shape(s, x, 2.65, 2.79, 2.78, fill)
        chip(s, f'0{i + 1}', x + .19, 2.85, .62, 'white', color)
        if i == 0: wheel(s, x + .32, 3.52, dirty=True)
        elif i == 1:
            for j in range(3): circle(s, x + .35 + j * .42, 3.66, .18, color)
            line(s, x + .34, 3.96, x + 1.5, 3.96, color, 2, True)
        elif i == 2:
            for j, h in enumerate([.22, .43, .64]): shape(s, x + .36 + j * .35, 4.04 - h, .23, h, color, kind=Shape.RECTANGLE)
        else: model(s, x + .42, 3.4, color=color)
        text(s, heading, x + .19, 4.25, 2.42, size=18, bold=True, color=color)
        text(s, body, x + .19, 4.73, 2.42, size=18)
        if i < 3: line(s, x + 2.84, 4.01, x + 3.02, 4.01, 'muted', 2, True)
    banner(s, 'В игре обучение начинается после передачи опыта роботу.', size=20)

    s = base(prs, 3, source[2])
    for i, (fill, color, title, body) in enumerate([
        ('mint', 'teal', 'Примеры', 'Покрытие + колёса до\nИзвестный результат'),
        ('lilac', 'purple', 'Модель', 'Находит правила,\nобъясняющие примеры'),
        ('yellow', 'ochre', 'Новый прогноз', 'Расход заряда\nСостояние колёс · проезд')]):
        x = .65 + i * 4.13
        shape(s, x, 2.6, 3.77, 2.98, fill)
        if i == 0:
            for j in range(3):
                shape(s, x + .34 + j * .12, 2.98 + j * .12, 1.04, .66, 'white', 'line')
                line(s, x + .52 + j * .12, 3.15 + j * .12, x + 1.09 + j * .12, 3.15 + j * .12, color, 2)
        elif i == 1: model(s, x + .58, 3.1, color=color)
        else:
            for px, py in [(x + .41, 3.39), (x + 1.1, 3.06), (x + 1.4, 3.64)]: circle(s, px, py, .19, color)
            line(s, x + .56, 3.45, x + 1.15, 3.15, color, 2)
            line(s, x + .56, 3.45, x + 1.46, 3.7, color, 2)
        text(s, title, x + .25, 4.13, 3.24, size=25, color=color, bold=True)
        text(s, body, x + .25, 4.68, 3.24, size=19)
        if i < 2: line(s, x + 3.82, 4.0, x + 4.07, 4.0, 'muted', 2, True)
    banner(s, 'На площадке известный ответ дают измерения поездки.', size=20)

    s = base(prs, 4, source[3])
    shape(s, .65, 2.62, 3.77, 2.94, 'white', 'line')
    text(s, 'Данные без меток', .94, 2.91, 3.19, size=24, bold=True)
    for row in range(3):
        for col in range(5): person(s, 1.08 + col * .59, 3.58 + row * .55, 'muted', .75)
    text(s, 'Поездки жителей', .94, 5.04, 3.14, size=18, color='muted')
    model(s, 5.63, 3.35, color='purple')
    line(s, 4.58, 3.9, 5.31, 3.9, 'muted', 2, True)
    line(s, 6.6, 3.9, 7.23, 3.9, 'muted', 2, True)
    text(s, 'Найти\nсходство', 4.64, 4.33, 2.73, size=22, color='purple', bold=True, align='center')
    for i, (fill, color) in enumerate([('mint', 'teal'), ('lilac', 'purple'), ('yellow', 'ochre')]):
        x = 7.42 + i * 1.78
        shape(s, x, 2.99, 1.58, 2.35, fill)
        for dx, dy in [(.29, .23), (.96, .23), (.625, .76), (.29, 1.29), (.96, 1.29)]:
            person(s, x + dx, 3.1 + dy, color, .68)
        text(s, f'Группа {i + 1}', x + .09, 4.82, 1.4, size=15, color=color, bold=True, align='center')
    banner(s, 'Пример: группы со схожими транспортными потребностями.', fill='lilac', color='purple', size=19)
    text(s, 'Другой метод обучения; в городской симуляции площадки не выполняется.', .7, 5.81, 11.9, size=16, color='muted', align='center')

    s = base(prs, 5, source[4])
    items = [(.65, 2.62, 'mint', 'teal', 'Агент выбирает действие', 'Например, куда двигаться.'),
             (7.02, 2.62, 'blue', 'steel', 'Среда отвечает', 'Действие меняет ситуацию.'),
             (7.02, 4.35, 'yellow', 'ochre', 'Агент получает награду', 'Учитывает успех и затраты.'),
             (.65, 4.35, 'lilac', 'purple', 'Обновляет стратегию', 'Учится выбирать следующие действия.')]
    for x, y, fill, color, heading, body in items:
        shape(s, x, y, 5.65, 1.26, fill)
        text(s, heading, x + .23, y + .18, 5.16, size=23, color=color, bold=True)
        text(s, body, x + .23, y + .76, 5.16, size=17)
    line(s, 6.4, 3.24, 6.91, 3.24, 'muted', 2, True)
    line(s, 9.84, 3.98, 9.84, 4.23, 'muted', 2, True)
    line(s, 6.91, 4.97, 6.4, 4.97, 'muted', 2, True)
    line(s, 3.47, 4.23, 3.47, 3.98, 'muted', 2, True)
    text(s, 'Другой метод обучения', .7, 5.86, 11.9, size=16, color='muted', align='center')
    banner(s, 'Робот площадки учится по измерениям; звёзды оценивают результат игры.', size=18)

    s = base(prs, 6, source[5])
    for x, fill in [(.65, 'mint'), (6.82, 'blue')]: shape(s, x, 2.57, 5.85, 3.12, fill)
    text(s, 'Робот · проверка модели', .95, 2.88, 5.22, size=25, color='teal', bold=True)
    line(s, 1.16, 3.85, 2.23, 3.45, 'teal', 2)
    line(s, 2.23, 3.45, 3.41, 3.9, 'teal', 2)
    for x, y in [(1.09, 3.78), (2.15, 3.38), (3.34, 3.82)]: circle(s, x, y, .18, 'teal')
    chip(s, 'НОВЫЙ ПУТЬ', 4.16, 3.56, 1.74)
    text(s, 'Прогноз → поездка → факт', .97, 4.37, 5.16, size=22, bold=True)
    text(s, 'Новые сочетания показывают,\nгде данных ещё не хватает.', .97, 4.94, 5.15, size=19)
    text(s, 'Город · агентная симуляция', 7.1, 2.88, 5.22, size=25, color='steel', bold=True)
    home(s, 7.41, 3.51, 'steel', .88)
    person(s, 8.6, 3.55, 'steel', 1)
    home(s, 9.53, 3.51, 'steel', .88)
    line(s, 8.11, 3.79, 8.48, 3.79, 'steel', 1.5, True)
    line(s, 8.99, 3.79, 9.4, 3.79, 'steel', 1.5, True)
    text(s, 'Жители выбирают действия', 7.12, 4.37, 5.15, size=21, bold=True)
    text(s, 'Правила заданы в модели.\nРешения меняют последствия.', 7.12, 4.94, 5.15, size=19)
    banner(s, 'Честное сравнение начинается с одинаковых исходных условий.', size=19)

    pptx = ROOT / 'ai-learning-introduction.pptx'; prs.save(pptx)
    (ROOT / 'presenter-notes.txt').write_text('\n\n'.join(f"СЛАЙД {i}: {v['title']}\n{v['notes']}" for i, v in enumerate(source, 1)), encoding='utf-8')
    return pptx, source


def build_guide(source):
    data = json.loads(source.read_text(encoding='utf-8'))
    target = ROOT / 'facilitator-script.json'; target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    doc = Document(); section = doc.sections[0]
    section.page_width = DocInches(8.27); section.page_height = DocInches(11.69)
    section.top_margin = section.bottom_margin = DocInches(.62)
    section.left_margin = section.right_margin = DocInches(.72)
    normal = doc.styles['Normal']; normal.font.name = FONT; normal.font.size = DocPt(11)
    normal.font.color.rgb = DocRGB.from_string(P['ink'])
    normal.paragraph_format.space_after = DocPt(6)
    normal.paragraph_format.line_spacing = 1.13
    for name, size, color in [('Title', 26, 'teal'), ('Heading 1', 19, 'teal'), ('Heading 2', 13, 'purple')]:
        style = doc.styles[name]; style.font.name = FONT; style.font.size = DocPt(size); style.font.bold = True
        style.font.color.rgb = DocRGB.from_string(P[color]); style.paragraph_format.space_after = DocPt(8)
    header = section.header.paragraphs[0]; header.text = 'поИИграем?  /  Памятка ведущего'
    header.style = normal; header.runs[0].font.size = DocPt(9); header.runs[0].font.color.rgb = DocRGB.from_string(P['muted'])
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    run = footer.add_run('Страница '); run.font.size = DocPt(9)
    fld = DocxElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    labels = {'say': 'Сказать', 'action': 'Сделать', 'question': 'Спросить', 'note': 'Ведущему'}
    text_lines = [data['title'], data['subtitle']]
    for n, page in enumerate(data['pages']):
        if n:
            doc.add_page_break()
        else:
            doc.add_paragraph(data['title'], 'Title'); doc.add_paragraph(data['subtitle'])
        doc.add_paragraph(page['title'], 'Heading 1'); text_lines.extend(['', page['title']])
        for block in page['blocks']:
            kind, message = block['kind'], block['text']; prefix = labels.get(kind)
            if kind == 'h2':
                doc.add_paragraph(message, 'Heading 2')
            elif kind == 'bullet':
                doc.add_paragraph(message, 'List Bullet')
            else:
                para = doc.add_paragraph()
                if prefix:
                    r = para.add_run(prefix + '. '); r.bold = True; r.font.color.rgb = DocRGB.from_string(P['teal' if kind in ['action', 'say'] else 'purple'])
                para.add_run(message)
            text_lines.append((prefix + '. ' if prefix else '') + message)
    path = ROOT / 'facilitator-guide.docx'; doc.save(path)
    (ROOT / 'facilitator-guide.txt').write_text('\n\n'.join(text_lines) + '\n', encoding='utf-8')
    return path


def export_pdf(paths):
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('LibreOffice is required to export the presentation and guide to PDF')
    with tempfile.TemporaryDirectory(prefix='ai-light-render-') as folder:
        output = Path(folder) / 'pdf'; output.mkdir(); profile = Path(folder) / 'profile'
        result = subprocess.run([executable, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(output), *map(str, paths)], text=True, capture_output=True, timeout=120)
        if result.returncode:
            raise RuntimeError(result.stderr + result.stdout)
        for path in paths:
            generated = output / path.with_suffix('.pdf').name
            if not generated.exists():
                raise RuntimeError('Missing PDF: ' + str(generated) + '\n' + result.stdout + result.stderr)
            shutil.copy2(generated, ROOT / generated.name)


PLAYER = '''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ИИ: данные, обучение и проверка</title><style>html,body{margin:0;background:#fcfbf8;width:100%;height:100%;overflow:hidden;font-family:Arial,sans-serif}main{height:calc(100% - 66px);display:flex;align-items:center;justify-content:center}img{width:100%;height:100%;object-fit:contain;display:block}nav{height:66px;box-sizing:border-box;display:flex;align-items:center;justify-content:space-between;padding:8px 24px;border-top:1px solid #d6e2dc;color:#233d40;background:#f3f7f1}.controls{display:flex;gap:8px;align-items:center}button,a{background:#fff;color:#233d40;border:1px solid #a6c7b9;border-radius:8px;font:16px Arial,sans-serif;padding:9px 13px;cursor:pointer;text-decoration:none}button:disabled{opacity:.35;cursor:default}button:focus-visible,a:focus-visible{outline:3px solid #307a71;outline-offset:2px}#counter{min-width:50px;text-align:center;font-size:14px}#downloadPanel{border:1px solid #a6c7b9;border-radius:14px;color:#233d40;background:#fcfbf8;padding:28px;max-width:600px}#downloadPanel::backdrop{background:#233d4070}.files{display:grid;gap:12px}#downloadClose{float:right}.hint{font-size:13px;color:#657b7c}@media(max-width:800px){nav{padding:8px}.hint{display:none}button,a{font-size:13px;padding:9px}}</style></head><body><main><img id="slide" src="slide-01.png" alt="Слайд 1 из 6"></main><nav><button id="downloads">Скачать материалы</button><span class="hint">← → переключить · F полный экран</span><div class="controls"><button id="prev" aria-label="Предыдущий слайд">←</button><span id="counter"></span><button id="next" aria-label="Следующий слайд">→</button><button id="full" aria-label="Полный экран">⛶</button></div></nav><dialog id="downloadPanel"><button id="downloadClose">Закрыть</button><h2>Материалы площадки</h2><div class="files"><a href="ai-learning-introduction.pptx" download>Презентация PowerPoint · 6 слайдов</a><a href="ai-learning-introduction.pdf" download>Презентация PDF</a><a href="facilitator-guide.docx" download>Текст ведущего с действиями · Word</a><a href="facilitator-guide.pdf" download>Текст ведущего · PDF</a><a href="facilitator-guide.txt" download>Текст ведущего · обычный текст</a></div></dialog><script>let n=1;const $=id=>document.getElementById(id);function show(v){n=Math.min(6,Math.max(1,v));$('slide').src='slide-'+String(n).padStart(2,'0')+'.png';$('slide').alt='Слайд '+n+' из 6';$('counter').textContent=n+' / 6';$('prev').disabled=n===1;$('next').disabled=n===6;}$('prev').onclick=()=>show(n-1);$('next').onclick=()=>show(n+1);async function full(){try{if(document.fullscreenElement)await document.exitFullscreen();else await document.documentElement.requestFullscreen();}catch{}}$('full').onclick=full;$('downloads').onclick=()=>$('downloadPanel').showModal();$('downloadClose').onclick=()=>$('downloadPanel').close();addEventListener('keydown',e=>{if($('downloadPanel').open)return;if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(n+1);}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(n-1);}if(e.key==='Home')show(1);if(e.key==='End')show(6);if(e.key.toLowerCase()==='f')full();});show(1);</script></body></html>'''


def finish_exports(source):
    pdf_path = ROOT / 'ai-learning-introduction.pdf'
    with fitz.open(pdf_path) as pdf:
        assert len(pdf) == 6
        for i, page in enumerate(pdf, 1):
            assert source[i - 1]['title'].replace(' ', '') in page.get_text().replace(' ', '').replace('\n', '')
            # QHD images keep the browser version sharp on the 27-inch monitors.
            page.get_pixmap(matrix=fitz.Matrix(8 / 3, 8 / 3), alpha=False).save(ROOT / f'slide-{i:02d}.png')
    (ROOT / 'index.html').write_text(PLAYER, encoding='utf-8')
    files = [ROOT / 'README.txt', ROOT / 'index.html', *sorted(ROOT.glob('slide-??.png')), ROOT / 'ai-learning-introduction.pptx', ROOT / 'ai-learning-introduction.pdf', ROOT / 'facilitator-guide.docx', ROOT / 'facilitator-guide.pdf', ROOT / 'facilitator-guide.txt', ROOT / 'presenter-notes.txt']
    with zipfile.ZipFile(ROOT / 'slides-and-browser-player.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        for file in files: archive.write(file, arcname=file.name)
    prs = Presentation(ROOT / 'ai-learning-introduction.pptx')
    assert len(prs.slides) == 6 and all(s.notes_slide.notes_text_frame.text for s in prs.slides)
    assert all(any(sh.has_text_frame for sh in s.shapes) for s in prs.slides), 'All slide text is editable'
    info = {'slides': 6, 'guide_pages': len(fitz.open(ROOT / 'facilitator-guide.pdf')), 'files': [{'name': file.name, 'bytes': file.stat().st_size} for file in files]}
    (ROOT / 'export-info.json').write_text(json.dumps(info, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(info, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(); parser.add_argument('--guide-source', type=Path, default=ROOT / 'facilitator-script.json')
    args = parser.parse_args()
    font_root = Path('/usr/share/fonts/truetype/liberation')
    pdfmetrics.registerFont(TTFont('Metric', str(font_root / 'LiberationSans-Regular.ttf')))
    pdfmetrics.registerFont(TTFont('MetricBold', str(font_root / 'LiberationSans-Bold.ttf')))
    pptx, source = build_slides(); guide = build_guide(args.guide_source)
    export_pdf([pptx, guide]); finish_exports(source)


if __name__ == '__main__':
    main()
